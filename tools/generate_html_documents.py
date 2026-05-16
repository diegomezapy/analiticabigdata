from __future__ import annotations

import html
import json
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree as ET

import fitz


ROOT = Path(__file__).resolve().parents[1]
PUBLIC_BASE = "https://diegomezapy.github.io/analiticabigdata/"
COURSE_NAME = "Analítica de Big Data"
TABLE_PREFIX = "__ABD_TABLE__:"
PLAN_XLSX = Path(r"G:\Mi unidad\FACEN_Software_Estadistico\Materiales_v2\Analitica_Big_Data\planificacion\FACEN_BigData_GD_1erPer_2026.xlsx")
COURSE_CONTEXT = {
    "Carrera": "Matemática Estadística",
    "Modalidad": "Virtual",
    "Año académico": "2026",
    "Periodo": "Primer período",
    "Institución": "FACEN · Universidad Nacional de Asunción",
}

UNIT_BRIEFS = {
    1: {
        "competencia": "Reconoce las características de Big Data y configura un entorno reproducible de trabajo con R/RStudio.",
        "ruta": [
            "Leer la orientación de la unidad y ubicar los conceptos de volumen, variedad y velocidad.",
            "Configurar el entorno de análisis y ejecutar operaciones básicas con datos tabulares.",
            "Elaborar una primera lectura exploratoria con tablas, resúmenes y gráficos simples.",
        ],
        "evidencia": "Bitácora breve con comandos utilizados, capturas de salida y reflexión sobre calidad de datos.",
    },
    2: {
        "competencia": "Aplica técnicas de preparación, visualización y análisis exploratorio para datos estructurados.",
        "ruta": [
            "Identificar variables, tipos de datos y valores faltantes.",
            "Construir visualizaciones que respondan preguntas analíticas concretas.",
            "Preparar variables mediante binarización, imputación, estandarización o escalado según corresponda.",
        ],
        "evidencia": "Reporte exploratorio con gráficos comentados y decisiones de preparación justificadas.",
    },
    3: {
        "competencia": "Implementa modelos predictivos básicos y compara resultados con criterios estadísticos.",
        "ruta": [
            "Definir variable objetivo, predictores y partición de datos.",
            "Entrenar modelos supervisados como KNN y regresión logística.",
            "Comparar desempeño e interpretar errores de clasificación.",
        ],
        "evidencia": "Cuaderno reproducible con modelo, matriz de confusión y discusión de resultados.",
    },
    4: {
        "competencia": "Interpreta modelos estadísticos y comunica hallazgos de forma pertinente para la toma de decisiones.",
        "ruta": [
            "Revisar métricas, supuestos y límites del modelo.",
            "Aplicar análisis discriminante lineal y contrastar resultados.",
            "Sintetizar conclusiones, riesgos y recomendaciones en un informe final.",
        ],
        "evidencia": "Informe final con problema, datos, metodología, resultados, limitaciones y recomendación.",
    },
}

ACTIVITY_BRIEFS = {
    1: {
        "tipo": "Actividad formativa",
        "proposito": "Asegurar lectura activa del material y preparación conceptual antes de las tareas aplicadas.",
        "producto": "Notas de lectura, respuestas guía y registro de dudas para tutoría.",
    },
    2: {
        "tipo": "Actividad interactiva",
        "proposito": "Contrastar ideas con el grupo mediante foro, caso o problema de aplicación.",
        "producto": "Intervención argumentada y réplica académica a un aporte de un compañero.",
    },
    3: {
        "tipo": "Actividad autónoma evaluable",
        "proposito": "Verificar comprensión individual mediante cuestionario o entrega breve.",
        "producto": "Cuestionario completado o evidencia de procedimiento según la consigna.",
    },
}


def unit_title(unit: int) -> str:
    return {
        1: "Introducción al manejo de datos en Big Data",
        2: "Visualización y exploración de datos",
        3: "Machine Learning",
        4: "Interpretación de resultados y modelos estadísticos",
    }[unit]


DOCUMENTS = [
    {
        "source": "guia_didactica/BigData_GuiaDid.pdf",
        "target": "guia_didactica/index.html",
        "download": "guia_didactica/BigData_GuiaDid.pdf",
        "title": "Guía Académica y Didáctica",
        "subtitle": "Analítica de Big Data",
        "kind": "Guía del curso",
    },
    *[
        {
            "source": f"unidades/unidad{unit}/Orientaciones_Unidad_Unidad{unit}.pdf",
            "target": f"unidades/unidad{unit}/orientaciones.html",
            "download": f"unidades/unidad{unit}/Orientaciones_Unidad_Unidad{unit}.pdf",
            "title": f"Orientaciones de la Unidad {unit}",
            "subtitle": unit_title(unit),
            "kind": "Orientaciones",
        }
        for unit in range(1, 5)
    ],
    *[
        {
            "source": f"unidades/unidad{unit}/Descripcion_Actividad_Unidad{unit}_Act{act}.pdf",
            "target": f"unidades/unidad{unit}/actividad-{act}.html",
            "download": f"unidades/unidad{unit}/Descripcion_Actividad_Unidad{unit}_Act{act}.pdf",
            "title": f"Descripción de Actividad {act}",
            "subtitle": f"Unidad {unit}: {unit_title(unit)}",
            "kind": "Actividad",
        }
        for unit in range(1, 5)
        for act in range(1, 4)
    ],
    {
        "source": "unidades/unidad1/Material_Lectura_Unidad1.pdf",
        "target": "unidades/unidad1/material-lectura.html",
        "download": "unidades/unidad1/Material_Lectura_Unidad1.pdf",
        "title": "Material de Lectura",
        "subtitle": "Unidad 1: Introducción al manejo de datos en Big Data",
        "kind": "Material de lectura",
    },
    *[
        {
            "source": f"unidades/unidad{unit}/Material_Lectura_Unidad{unit}_Extendido.docx",
            "target": f"unidades/unidad{unit}/material-extendido.html",
            "download": f"unidades/unidad{unit}/material-extendido.pdf",
            "title": "Material de Lectura Extendido",
            "subtitle": f"Unidad {unit}: {unit_title(unit)}",
            "kind": "Material extendido",
        }
        for unit in range(1, 5)
    ],
]


def load_plan_rows() -> list[dict[str, object]]:
    if not PLAN_XLSX.exists():
        return []
    try:
        import openpyxl
    except Exception:
        return []
    workbook = openpyxl.load_workbook(PLAN_XLSX, data_only=True)
    sheet = workbook["cronograma"]
    rows: list[dict[str, object]] = []
    for row in sheet.iter_rows(min_row=19, max_row=44, values_only=True):
        unit = row[0]
        if unit in (None, ""):
            continue
        try:
            unit_num = int(unit)
        except (TypeError, ValueError):
            continue
        start = format_plan_date(row[9])
        end = format_plan_date(row[10])
        if not start or not end:
            continue
        rows.append({
            "unit": unit_num,
            "unit_name": normalize(str(row[1] or "")),
            "unit_objective": normalize(str(row[2] or "")),
            "content": normalize(str(row[3] or "")),
            "activity": normalize(str(row[4] or "")),
            "type": normalize(str(row[5] or "")),
            "points": normalize(str(row[6] if row[6] is not None else "")),
            "activity_objective": normalize(str(row[7] or "")),
            "duration": normalize(str(row[8] if row[8] is not None else "")),
            "start": start,
            "end": end,
        })
    return rows


def format_plan_date(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        return ""
    text = str(value).strip()
    if not text or text == "00:00:00":
        return ""
    try:
        return datetime.fromisoformat(text).date().isoformat()
    except ValueError:
        return text


def format_plan_display_date(value: object) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    try:
        return datetime.fromisoformat(text).strftime("%d/%m/%Y")
    except ValueError:
        return text


def display_plan_cell(row: dict[str, object], key: str) -> str:
    value = row.get(key, "")
    if key in {"start", "end"}:
        return format_plan_display_date(value)
    return str(value)


def plan_rows_for_unit(unit: int) -> list[dict[str, object]]:
    return [row for row in PLAN_ROWS if row["unit"] == unit]


def plan_row_for_activity(unit: int, act: int) -> dict[str, object] | None:
    prefix = f"{act}."
    for row in plan_rows_for_unit(unit):
        if str(row["activity"]).strip().startswith(prefix):
            return row
    return None


def render_plan_table(rows: list[dict[str, object]], columns: list[tuple[str, str]], compact: bool = True) -> str:
    if not rows:
        return ""
    header = "".join(f"<th>{html.escape(label)}</th>" for _, label in columns)
    body_rows = []
    for row in rows:
        cells = "".join(f"<td>{html.escape(display_plan_cell(row, key))}</td>" for key, _ in columns)
        body_rows.append(f"<tr>{cells}</tr>")
    class_name = "doc-table doc-table-compact" if compact else "doc-table"
    return f'<div class="doc-table-wrap"><table class="{class_name}"><tr>{header}</tr>{"".join(body_rows)}</table></div>'


def render_planning_frontmatter(spec: dict[str, str]) -> str:
    unit = spec_unit(spec)
    act = spec_activity(spec)
    if spec["kind"] == "Guía del curso":
        schedule_rows = [row for row in PLAN_ROWS if row["unit"] in {1, 2, 3, 4}]
        exam_rows = [row for row in PLAN_ROWS if row["unit"] == 5]
        tutoring_rows = [row for row in PLAN_ROWS if row["unit"] == 9]
        schedule = render_plan_table(schedule_rows, [
            ("unit", "Unidad"),
            ("content", "Contenido"),
            ("activity", "Actividad"),
            ("type", "Tipo"),
            ("points", "Puntaje"),
            ("start", "Inicio"),
            ("end", "Fin"),
        ])
        exams = render_plan_table(exam_rows, [
            ("activity", "Evaluación"),
            ("type", "Tipo"),
            ("points", "Puntaje"),
            ("start", "Inicio"),
            ("end", "Fin"),
        ])
        tutoring = render_plan_table(tutoring_rows, [
            ("activity", "Tutoría"),
            ("start", "Fecha"),
            ("type", "Tipo"),
        ])
        return f"""
    <section class="doc-study-note">
      <h2>Cronograma oficial de la planificación</h2>
      <p>La siguiente tabla proviene de la planilla de planificación del primer período 2026 y reemplaza los cuadros extraídos automáticamente del PDF.</p>
      {schedule}
    </section>
    <section class="doc-study-note">
      <h2>Evaluaciones parciales y finales</h2>
      {exams}
    </section>
    <section class="doc-study-note">
      <h2>Tutorías sincrónicas</h2>
      {tutoring}
    </section>
"""
    if unit and not act:
        schedule = render_plan_table(plan_rows_for_unit(unit), [
            ("content", "Contenido"),
            ("activity", "Actividad"),
            ("type", "Tipo"),
            ("points", "Puntaje"),
            ("duration", "Días"),
            ("start", "Inicio"),
            ("end", "Fin"),
        ])
        if schedule:
            return f"""
    <section class="doc-study-note">
      <h2>Actividades oficiales de la unidad</h2>
      {schedule}
    </section>
"""
    if unit and act:
        row = plan_row_for_activity(unit, act)
        if row:
            schedule = render_plan_table([row], [
                ("content", "Contenido"),
                ("activity", "Actividad"),
                ("type", "Tipo"),
                ("points", "Puntaje"),
                ("duration", "Días"),
                ("start", "Inicio"),
                ("end", "Fin"),
            ], compact=False)
            return f"""
    <section class="doc-study-note">
      <h2>Datos oficiales de la actividad</h2>
      {schedule}
    </section>
"""
    return ""


def extract_pdf(path: Path) -> list[str]:
    doc = fitz.open(path)
    lines: list[str] = []
    for page_index, page in enumerate(doc, start=1):
        if page_index > 1:
            lines.append(f"Página {page_index}")
        page_items: list[tuple[float, str]] = []
        table_rects: list[fitz.Rect] = []
        try:
            tables = page.find_tables().tables
        except Exception:
            tables = []
        for table in tables:
            rows = clean_table_rows(table.extract())
            if rows:
                table_rect = fitz.Rect(table.bbox)
                table_rects.append(table_rect)
                page_items.append((table_rect.y0, table_marker(rows)))
        for block in page.get_text("blocks"):
            block_rect = fitz.Rect(block[:4])
            if any(overlap_ratio(block_rect, table_rect) > 0.2 for table_rect in table_rects):
                continue
            raw_lines = [normalize(raw) for raw in block[4].splitlines()]
            raw_lines = [line for line in raw_lines if line]
            if not raw_lines:
                continue
            if len(raw_lines) > 1 and all(is_heading_or_meta(line) for line in raw_lines):
                page_items.extend((block_rect.y0, line) for line in raw_lines)
            else:
                page_items.append((block_rect.y0, normalize(" ".join(raw_lines))))
        lines.extend(item for _, item in sorted(page_items, key=lambda item: item[0]))
    return collapse_repeated(lines)


def extract_docx(path: Path) -> list[str]:
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    lines: list[str] = []
    with zipfile.ZipFile(path) as docx:
        root = ET.fromstring(docx.read("word/document.xml"))
    body = root.find("w:body", ns)
    if body is None:
        return lines
    for element in body:
        tag = strip_ns(element.tag)
        if tag == "p":
            text = paragraph_text(element, ns)
            if text:
                lines.append(text)
        elif tag == "tbl":
            rows = []
            for row in element.findall(".//w:tr", ns):
                cells = []
                for cell in row.findall("./w:tc", ns):
                    text = " ".join(
                        paragraph_text(p, ns)
                        for p in cell.findall(".//w:p", ns)
                    ).strip()
                    if text:
                        cells.append(text)
                if cells:
                    rows.append(cells)
            if rows:
                lines.append(table_marker(rows))
    return collapse_repeated(lines)


def table_marker(rows: list[list[str]]) -> str:
    return TABLE_PREFIX + json.dumps(rows, ensure_ascii=False)


def clean_table_rows(rows: list[list[str | None]]) -> list[list[str]]:
    cleaned: list[list[str]] = []
    for row in rows:
        cells: list[str] = []
        previous = ""
        for cell in row:
            text = normalize(str(cell or ""))
            if not text or text == previous:
                continue
            cells.append(text)
            previous = text
        if cells and cells not in cleaned:
            cleaned.append(cells)
    return cleaned


def overlap_ratio(a: fitz.Rect, b: fitz.Rect) -> float:
    intersection = a & b
    if intersection.is_empty or a.get_area() == 0:
        return 0.0
    return intersection.get_area() / a.get_area()


def paragraph_text(element: ET.Element, ns: dict[str, str]) -> str:
    parts = []
    for text_node in element.findall(".//w:t", ns):
        if text_node.text:
            parts.append(text_node.text)
    return normalize("".join(parts))


def strip_ns(tag: str) -> str:
    return tag.split("}", 1)[-1]


def normalize(value: str) -> str:
    value = re.sub(r"\s+", " ", value.replace("\xa0", " ")).strip()
    return value


def collapse_repeated(lines: list[str]) -> list[str]:
    collapsed: list[str] = []
    previous = None
    for line in lines:
        if line == previous:
            continue
        collapsed.append(line)
        previous = line
    return collapsed


PLAN_ROWS = load_plan_rows()


def line_to_html(line: str, index: int) -> str:
    if line.startswith(TABLE_PREFIX):
        return render_table_html(line)
    escaped = html.escape(line)
    if not line:
        return ""
    if index == 0:
        return f"<p>{escaped}</p>"
    if re.match(r"^(Página \d+|pág\. \d+|www\.virtual\.facen\.una\.py)", line, re.I):
        return f'<p class="doc-meta-line">{escaped}</p>'
    if re.match(r"^•\s+", line):
        return f'<p class="doc-bullet">{escaped}</p>'
    if " | " in line:
        return render_table_html(table_marker([line.split(" | ")]))
    if is_equation(line):
        return f'<div class="doc-equation">{escaped}</div>'
    if is_heading(line):
        return f"<h2>{escaped}</h2>"
    return f"<p>{escaped}</p>"


def filter_source_lines(lines: list[str], spec: dict[str, str]) -> list[str]:
    filtered: list[str] = []
    skip_next_meta = False
    for line in lines:
        if line.startswith(TABLE_PREFIX) and is_planning_table_marker(line):
            continue
        if is_planning_noise_line(line):
            skip_next_meta = True
            continue
        if skip_next_meta and re.match(r"^(Referencias: AA|AA: Actividad|AI: Actividad|AF: Actividad)", line, re.I):
            skip_next_meta = False
            continue
        filtered.append(line)
    return collapse_repeated(filtered)


def is_planning_noise_line(line: str) -> bool:
    patterns = [
        r"^5\.\s*Cronograma\b",
        r"^Unidad\s+Nro\b",
        r"^Unidad\s+Descripción\b",
        r"^Actividad\s+Descripción\s+Tipo",
        r"^Puntaje\s+Fecha\s+inicio",
        r"^\d\.0\s+",
        r"^(Primer|Segundo)\s+(parcial|final)$",
    ]
    return any(re.search(pattern, line, re.I) for pattern in patterns)


def is_planning_table_marker(marker: str) -> bool:
    try:
        rows = json.loads(marker.removeprefix(TABLE_PREFIX))
    except json.JSONDecodeError:
        return False
    text = " ".join(" ".join(str(cell) for cell in row) for row in rows)
    signals = [
        "Fecha inicio",
        "Fecha fin",
        "Actividad Descripción",
        "Tipo act",
        "Unidad Descripción",
    ]
    return sum(1 for signal in signals if signal.lower() in text.lower()) >= 2


def render_table_html(marker: str) -> str:
    try:
        rows = json.loads(marker.removeprefix(TABLE_PREFIX))
    except json.JSONDecodeError:
        return ""
    if looks_like_code_rows(rows):
        return render_code_html(rows)
    html_rows = []
    for row_index, row in enumerate(rows):
        tag = "th" if row_index == 0 else "td"
        cells = "".join(f"<{tag}>{html.escape(str(cell))}</{tag}>" for cell in row)
        html_rows.append(f"<tr>{cells}</tr>")
    return f'<div class="doc-table-wrap"><table class="doc-table">{"".join(html_rows)}</table></div>'


def looks_like_code_rows(rows: list[list[str]]) -> bool:
    if not rows:
        return False
    max_cols = max(len(row) for row in rows)
    text = "\n".join(" ".join(str(cell) for cell in row) for row in rows)
    code_tokens = [
        "library(", "import ", "from ", "<-", "#", "print(", "head(",
        "summary(", "install.packages", "read.csv", "read_csv", "pd.",
        "pl.", "ggplot(", "sns.", "plt.", "```", "knitr::", "set.seed(",
        "data.frame(", "function(", "confusionMatrix(", "fread(",
        "hist(", "plot(", "boxplot(", "abline(", "mean(", "sd(",
        "getwd(", "list.files(",
    ]
    tree_tokens = ["├", "└", "│", ".Rproj", "/"]
    code_score = sum(1 for token in code_tokens if token in text)
    if max_cols == 1 and code_score >= 2:
        return True
    if max_cols == 1 and len(rows) >= 2 and str(rows[0][0]).lstrip().startswith("#"):
        return True
    if max_cols == 1 and len(rows) >= 3 and any(token in text for token in tree_tokens):
        return True
    return False


def render_code_html(rows: list[list[str]]) -> str:
    lines: list[str] = []
    for row in rows:
        joined = " ".join(str(cell) for cell in row).strip()
        if joined:
            lines.extend(split_compacted_code_line(joined))
    code = "\n".join(line for line in lines if line.strip())
    return f'<pre class="doc-code"><code>{html.escape(code)}</code></pre>'


def split_compacted_code_line(line: str) -> list[str]:
    if len(line) < 90:
        return [line]
    attached_starts = [
        r"df\s*=", r"df\s*<-", r"df_", r"datos\s*<-", r"ggplot\(", r"ggcorr\(",
        r"sns\.", r"plt\.", r"print\(", r"head\(", r"summary\(",
        r"corr\s*=", r"imputer\s*=", r"trainIndex\b", r"model\s*<-",
        r"predictions\s*<-", r"hist\(", r"plot\(", r"boxplot\(",
    ]
    formatted = line
    for start in attached_starts:
        formatted = re.sub(r"(# [^\n#]*?)(?=" + start + ")", r"\1\n", formatted)
    break_patterns = [
        r"(?<!^)(?=import\s)",
        r"(?<!^)(?=from\s)",
        r"(?<!^)(?=library\()",
        r"(?<!^)(?=install\.packages\()",
        r"(?<!^)(?=#\s)",
        r"(?<!^)(?=data\()",
        r"(?<!^)(?=set\.seed\()",
        r"(?<!^)(?=trainIndex\b)",
        r"(?<!^)(?=train\s*<-)",
        r"(?<!^)(?=test\s*<-)",
        r"(?<!^)(?=model\s*<-)",
        r"(?<!^)(?=predictions\s*<-)",
        r"(?<!^)(?=normalize\s*<-)",
        r"(?<!^)(?=iris_norm\s*<-)",
        r"(?<!^)(?=imputer\s*=)",
        r"(?<!^)(?=df(?:_|\s*=|\s*<-))",
        r"(?<!^)(?=datos(?:_|\s*<-))",
        r"(?<!^)(?=corr\s*=)",
        r"(?<![\w.])(?=print\()",
        r"(?<![\w.])(?=head\()",
        r"(?<![\w.])(?=summary\()",
        r"(?<![\w.])(?=confusionMatrix\()",
        r"(?<![\w.])(?=plot\()",
        r"(?<!^)(?=ggplot\()",
        r"(?<!^)(?=sns\.)",
        r"(?<!^)(?=plt\.)",
        r"(?<![\w.])(?=dim\()",
        r"(?<![\w.])(?=nrow\()",
        r"(?<![\w.])(?=ncol\()",
        r"(?<![\w.])(?=colnames\()",
        r"(?<![\w.])(?=str\()",
        r"(?<![\w.])(?=glimpse\()",
        r"(?<![\w.])(?=is\.na\()",
        r"(?<![\w.])(?=colSums\()",
        r"(?<![\w.])(?=sum\()",
        r"(?<![\w.])(?=unique\()",
        r"(?<![\w.])(?=table\()",
    ]
    for pattern in break_patterns:
        formatted = re.sub(pattern, "\n", formatted)
    return [part.strip() for part in formatted.splitlines()]


def is_equation(line: str) -> bool:
    if len(line) > 180 or is_heading(line):
        return False
    math_tokens = [
        "=", "≤", "≥", "≠", "≈", "∑", "√", "β", "α", "λ", "μ", "σ",
        "logit", "Pr(", "P(", "y =", "Y =", "f(", "^2", "\\frac"
    ]
    if any(token in line for token in math_tokens):
        if re.search(r"\b(R|Python|Código|Unidad|Fecha|Nombre|Carrera)\b", line, re.I):
            return False
        return True
    return False


def is_heading_or_meta(line: str) -> bool:
    return bool(re.match(r"^(Página \d+|pág\. \d+|www\.virtual\.facen\.una\.py)", line, re.I)) or is_heading(line)


def is_heading(line: str) -> bool:
    heading_patterns = [
        r"^\d+(\.\d+)*\.?\s+\S",
        r"^Unidad\s+\d",
        r"^Departamento de Educación",
        r"^Analítica de Big Data$",
        r"^Guía Didáctica$",
        r"^Orientaciones$",
        r"^Material de Lectura",
        r"^Descripción de Actividad",
        r"^Datos generales",
        r"^Datos del profesor",
        r"^Objetivo",
        r"^Actividades",
        r"^Metodología",
        r"^Contenido",
        r"^Cronograma",
        r"^Evaluación",
        r"^Referencias",
        r"^Tabla\s+\d",
        r"^Figura\s+\d",
        r"^Ejemplo",
        r"^(Python|R) Code:$",
    ]
    return any(re.match(pattern, line, re.I) for pattern in heading_patterns)


def render_document(spec: dict[str, str], lines: list[str]) -> str:
    target = spec["target"]
    depth = len(Path(target).parts) - 1
    css_path = "../" * depth + "css/documentos.css"
    asset_path = "../" * depth + "assets/academic/"
    canonical = PUBLIC_BASE + target.replace("\\", "/")
    download = spec.get("download", "")
    download_link = ""
    if download:
        download_href = "../" * depth + download
        download_link = f'<a class="doc-download" href="{download_href}" download>Descargar PDF</a>'
    frontmatter = render_academic_frontmatter(spec, asset_path)
    planning = render_planning_frontmatter(spec)
    body_lines = filter_source_lines(lines, spec)
    body = "\n".join(line_to_html(line, i) for i, line in enumerate(body_lines))
    return f"""<!DOCTYPE html>
<html lang="es">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>{html.escape(spec["title"])} — Analítica de Big Data</title>
  <meta name="description" content="{html.escape(spec["kind"])} de Analítica de Big Data FACEN.">
  <link rel="canonical" href="{canonical}">
  <link rel="stylesheet" href="{css_path}">
  <link rel="manifest" href="{ '../' * depth }manifest.json">
  <meta name="theme-color" content="#f8fafc">
</head>
<body>
  <main class="doc-shell">
    <div class="doc-brand-strip">
      <img src="{asset_path}facen-header.png" alt="FACEN Universidad Nacional de Asunción">
    </div>
    <header class="doc-header">
      <div class="doc-header-copy">
        <div class="doc-kicker">{html.escape(spec["kind"])}</div>
        <h1>{html.escape(spec["title"])}</h1>
        <p>{html.escape(spec["subtitle"])}</p>
      </div>
      <div class="doc-actions">
        <a class="doc-home" href="{ '../' * depth }dashboard.html">Volver a la plataforma</a>
        {download_link}
      </div>
    </header>
{frontmatter}
{planning}
    <article class="doc-content">
      <h2>Contenido del documento fuente</h2>
      {body}
    </article>
    <footer class="doc-footer">
      <img src="{asset_path}facen-footer.png" alt="Educación a Distancia FACEN UNA">
      <div>www.virtual.facen.una.py · Dpto. de Educación a Distancia</div>
    </footer>
  </main>
</body>
</html>
"""


def spec_unit(spec: dict[str, str]) -> int | None:
    match = re.search(r"unidad(\d)", spec.get("target", ""), re.I)
    if match:
        return int(match.group(1))
    return None


def spec_activity(spec: dict[str, str]) -> int | None:
    match = re.search(r"actividad-(\d)", spec.get("target", ""), re.I)
    if match:
        return int(match.group(1))
    return None


def render_academic_frontmatter(spec: dict[str, str], asset_path: str) -> str:
    unit = spec_unit(spec)
    act = spec_activity(spec)
    meta = "".join(
        f"<div><span>{html.escape(key)}</span><strong>{html.escape(value)}</strong></div>"
        for key, value in COURSE_CONTEXT.items()
    )
    if spec["kind"] == "Guía del curso":
        return f"""
    <section class="doc-study-note">
      <h2>Guía académica de la asignatura</h2>
      <p>Documento marco para comprender la organización pedagógica, los objetivos, la metodología, el cronograma y los criterios de evaluación de la asignatura.</p>
    </section>
    <section class="doc-meta-grid">{meta}</section>
    <section class="doc-study-note">
      <h2>Uso recomendado</h2>
      <p>Revise esta guía antes de iniciar las unidades. Utilícela como contrato didáctico: allí se indican las expectativas de participación, la secuencia de actividades, los plazos y la forma en que se consolidan las evidencias de aprendizaje.</p>
    </section>
"""
    if unit:
        brief = UNIT_BRIEFS[unit]
        ruta = "".join(f"<li>{html.escape(item)}</li>" for item in brief["ruta"])
        activity_html = ""
        if act:
            activity = ACTIVITY_BRIEFS[act]
            activity_html = f"""
      <div class="doc-activity-brief">
        <h3>{html.escape(activity["tipo"])}</h3>
        <p><strong>Propósito:</strong> {html.escape(activity["proposito"])}</p>
        <p><strong>Producto esperado:</strong> {html.escape(activity["producto"])}</p>
      </div>
"""
        return f"""
    <section class="doc-unit-card no-image">
      <div class="doc-unit-number">Unidad {unit}</div>
      <div>
        <div class="doc-kicker">Analítica de Big Data</div>
        <h2>{html.escape(unit_title(unit))}</h2>
        <p>{html.escape(brief["competencia"])}</p>
      </div>
    </section>
    <section class="doc-meta-grid">{meta}<div><span>Unidad</span><strong>{unit}</strong></div></section>
    <section class="doc-study-note">
      <h2>Ruta de trabajo sugerida</h2>
      <ol>{ruta}</ol>
      <p><strong>Evidencia de aprendizaje:</strong> {html.escape(brief["evidencia"])}</p>
      {activity_html}
    </section>
"""
    return f'<section class="doc-meta-grid">{meta}</section>'


def render_index() -> str:
    document_items = "\n".join(
        f'<a href="{html.escape(spec["target"])}"><span>{html.escape(spec["kind"])}</span>{html.escape(spec["title"])} — {html.escape(spec["subtitle"])}</a>'
        for spec in DOCUMENTS
    )
    resource_items = """
<a href="guia_didactica/BigData_GuiaDid.pdf" download><span>PDF descargable</span>Guía Académica y Didáctica — PDF</a>
<a href="practicas/index.html"><span>Prácticas</span>Laboratorios guiados con datos CSV y código R</a>
<a href="laboratorio/index.html"><span>Laboratorio ejecutable</span>R con data.table y Python con Polars en el navegador</a>
<a href="data/datasets/clientes_compras.csv" download><span>Datos CSV</span>Clientes y compras — práctica Unidad 1</a>
<a href="data/datasets/estudiantes_ead.csv" download><span>Datos CSV</span>Estudiantes EAD — práctica Unidad 2</a>
<a href="data/datasets/modelos_credito.csv" download><span>Datos CSV</span>Modelos de crédito — prácticas Unidades 3 y 4</a>
"""
    items = resource_items + document_items
    return f"""<!DOCTYPE html>
<html lang="es">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>Documentos HTML — Analítica de Big Data</title>
  <meta name="description" content="Índice de documentos HTML insertables de Analítica de Big Data FACEN.">
  <link rel="canonical" href="{PUBLIC_BASE}documentos.html">
  <link rel="stylesheet" href="css/documentos.css">
  <link rel="manifest" href="manifest.json">
  <meta name="theme-color" content="#f8fafc">
</head>
<body>
  <main class="doc-shell">
    <div class="doc-brand-strip">
      <img src="assets/academic/facen-header.png" alt="FACEN Universidad Nacional de Asunción">
    </div>
    <header class="doc-header">
      <div class="doc-kicker">Documentos HTML</div>
      <h1>Analítica de Big Data</h1>
      <p>Guía, orientaciones, actividades y materiales listos para abrir o insertar desde GitHub Pages.</p>
      <a class="doc-home" href="dashboard.html">Volver a la plataforma</a>
    </header>
    <section class="doc-meta-grid">
      {"".join(f"<div><span>{html.escape(key)}</span><strong>{html.escape(value)}</strong></div>" for key, value in COURSE_CONTEXT.items())}
    </section>
    <nav class="doc-index">
      {items}
    </nav>
    <footer class="doc-footer">
      <img src="assets/academic/facen-footer.png" alt="Educación a Distancia FACEN UNA">
      <div>www.virtual.facen.una.py · Dpto. de Educación a Distancia</div>
    </footer>
  </main>
</body>
</html>
"""


def main() -> None:
    for spec in DOCUMENTS:
        source = ROOT / spec["source"]
        target = ROOT / spec["target"]
        if source.suffix.lower() == ".pdf":
            lines = extract_pdf(source)
        elif source.suffix.lower() == ".docx":
            lines = extract_docx(source)
        else:
            raise ValueError(f"Unsupported source type: {source}")
        download = spec.get("download")
        if download and source.suffix.lower() == ".docx":
            write_simple_pdf(ROOT / download, spec, lines)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(render_document(spec, lines), encoding="utf-8", newline="\n")
        print(f"generated {target.relative_to(ROOT)}")
    (ROOT / "documentos.html").write_text(render_index(), encoding="utf-8", newline="\n")
    print("generated documentos.html")


def write_simple_pdf(target: Path, spec: dict[str, str], lines: list[str]) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    doc = fitz.open()
    header_img = ROOT / "assets/academic/facen-header.png"
    footer_img = ROOT / "assets/academic/facen-footer.png"
    header_xref = 0
    footer_xref = 0

    def decorate(page: fitz.Page) -> None:
        nonlocal header_xref, footer_xref
        if header_img.exists():
            header_xref = page.insert_image(
                fitz.Rect(320, 0, 595, 67),
                filename=str(header_img) if header_xref == 0 else None,
                xref=header_xref,
                keep_proportion=True
            )
        if footer_img.exists():
            footer_xref = page.insert_image(
                fitz.Rect(0, 797, 595, 842),
                filename=str(footer_img) if footer_xref == 0 else None,
                xref=footer_xref,
                keep_proportion=False
            )
        page.insert_text((54, 812), "Educación a Distancia FACEN UNA", fontsize=9, fontname="helv", color=(1, 1, 1))

    page = doc.new_page(width=595, height=842)
    decorate(page)
    margin = 54
    y = 96

    def add_page() -> fitz.Page:
        new_page = doc.new_page(width=595, height=842)
        decorate(new_page)
        return new_page

    def put(text: str, size: int = 10, bold: bool = False) -> None:
        nonlocal page, y
        font = "helv"
        max_width = page.rect.width - margin * 2
        for chunk in wrap_text(text, max_chars=max(45, int(max_width / (size * 0.48)))):
            if y > page.rect.height - margin:
                page = add_page()
                y = 96
            page.insert_text((margin, y), chunk, fontsize=size, fontname=font, color=(0.08, 0.12, 0.2))
            y += size * (1.55 if bold else 1.45)
        y += 2

    put("FACEN · Universidad Nacional de Asunción", size=9)
    put(spec["title"], size=18, bold=True)
    put(spec["subtitle"], size=12)
    y += 8
    for line in lines:
        if line.startswith(TABLE_PREFIX):
            try:
                rows = json.loads(line.removeprefix(TABLE_PREFIX))
            except json.JSONDecodeError:
                rows = []
            for row in rows:
                put(" | ".join(str(cell) for cell in row), size=8)
            y += 4
            continue
        if not line:
            continue
        put(line, size=12 if is_heading(line) else 10, bold=is_heading(line))
    doc.save(target)
    doc.close()
    print(f"generated {target.relative_to(ROOT)}")


def wrap_text(text: str, max_chars: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current: list[str] = []
    length = 0
    for word in words:
        extra = len(word) + (1 if current else 0)
        if current and length + extra > max_chars:
            lines.append(" ".join(current))
            current = [word]
            length = len(word)
        else:
            current.append(word)
            length += extra
    if current:
        lines.append(" ".join(current))
    return lines or [text]


if __name__ == "__main__":
    main()
